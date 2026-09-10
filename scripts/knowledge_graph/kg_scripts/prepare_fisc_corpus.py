#!/usr/bin/env python3
"""Prepare FISC redacted/unredacted PDF pairs for the packaged KG pipeline.

This script starts directly from paired PDFs such as ``Pair_1-1.pdf`` and
``Pair_1-2.pdf`` under ``data/raw/fisc/``. It builds:

1. a transcription-style staging root with per-page text, aligned diffs, and
   pair metadata; and
2. a final pair-style ``prepared`` corpus root consumed by
   ``kg_scripts/run_pipeline.py``.

Each PDF is transcribed page-by-page with the packaged OpenAI image-based
transcription helper so the prepared corpus is generated through one consistent
path.
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

try:
    from tqdm.auto import tqdm
except Exception:  # pragma: no cover
    tqdm = None

ROOT = Path(__file__).resolve().parents[1]
PACKAGE_ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

from upstream_pair_processing.openai_pdf_pipeline import (
    PageBlock,
    PDFPair,
    _json_safe,
    build_difference_artifacts,
    ensure_openai_client,
    global_align_pages,
    save_split_pages,
    split_transcription_into_pages,
    transcribe_pdf_with_openai,
)
from upstream_pair_processing.postprocess_redaction_labels import run_postprocess


FISC_TRANSCRIPTION_PROMPT = """You are a careful OCR transcription assistant working on PUBLICLY RELEASED court records.

These PDFs are public documents from open releases such as the Senate Judiciary Committee, FBI Vault, CourtListener, or other public court archives. They are not secret source material for you to summarize, analyze, or refuse. Your task is only to transcribe the visible page text faithfully.

Please follow these rules exactly:

1. Transcribe page by page into plain text.
2. Start each page with [PAGE XX], using two digits.
3. Transcribe only visible text that appears on the page.
4. If text is redacted, blacked out, blanked, or removed, skip that hidden text silently. Do not guess it.
5. Do not summarize, explain, warn, refuse, or comment on the document.
6. Do not wrap the answer in markdown code fences.
7. Ignore scanner borders, image noise, specks, and non-text artifacts.
8. Keep headings, docket labels, paragraph order, and list structure when visible.
9. If a page contains no readable text, output only its [PAGE XX] label.
10. Return plain text only.

Transcribe the provided PDF now."""


PAIR_FILENAME_RE = re.compile(r"^Pair_(\d+)-([12])\.pdf$", re.IGNORECASE)
TOKEN_RE = re.compile(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*")
YEAR_MONTH_RE = re.compile(
    r"\b(20\d{2})\s+"
    r"(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|SEPT|OCT|NOV|DEC)\b",
    re.IGNORECASE,
)
MONTH_YEAR_RE = re.compile(
    r"\b("
    r"JAN(?:UARY)?|FEB(?:RUARY)?|MAR(?:CH)?|APR(?:IL)?|MAY|JUN(?:E)?|"
    r"JUL(?:Y)?|AUG(?:UST)?|SEP(?:T(?:EMBER)?)?|OCT(?:OBER)?|NOV(?:EMBER)?|DEC(?:EMBER)?"
    r")\s+(20\d{2})\b",
    re.IGNORECASE,
)
FULL_DATE_RE = re.compile(
    r"\b("
    r"JAN(?:UARY)?|FEB(?:RUARY)?|MAR(?:CH)?|APR(?:IL)?|MAY|JUN(?:E)?|"
    r"JUL(?:Y)?|AUG(?:UST)?|SEP(?:T(?:EMBER)?)?|OCT(?:OBER)?|NOV(?:EMBER)?|DEC(?:EMBER)?"
    r")\s+(\d{1,2}),\s*(20\d{2})\b",
    re.IGNORECASE,
)
DAY_MONTH_YEAR_RE = re.compile(
    r"\b(\d{1,2})\s+("
    r"JAN(?:UARY)?|FEB(?:RUARY)?|MAR(?:CH)?|APR(?:IL)?|MAY|JUN(?:E)?|"
    r"JUL(?:Y)?|AUG(?:UST)?|SEP(?:T(?:EMBER)?)?|OCT(?:OBER)?|NOV(?:EMBER)?|DEC(?:EMBER)?"
    r")\s+(20\d{2})\b",
    re.IGNORECASE,
)
FURNITURE_PATTERNS = [
    re.compile(r"^SENATE-FISA\d+(?:-\d+)?$", re.IGNORECASE),
    re.compile(r"^Case\s+\S+\s+Document\s+\d+\s+Filed\b.*\bPage\s+\d+\s+of\s+\d+$", re.IGNORECASE),
    re.compile(r"^(?:b\d+[A-Z]?-\d+[,\s]*)+$", re.IGNORECASE),
]
CLASSIFICATION_LINE_RE = re.compile(r"^(?:TOP\s+SECRET(?:\/\/NOFORN)?|SECRET|CONFIDENTIAL|UNCLASSIFIED)\s*$", re.IGNORECASE)
HEADER_NOISE_PREFIX_RE = re.compile(r"^(?:Classify By|Reason|Derived From|Declassify On)\s*:", re.IGNORECASE)
BOILERPLATE_KEYWORDS = {
    "clerk of court",
    "classify by",
    "classified by",
    "declassify on",
    "declassified by",
    "derived from",
    "foreign intelligence surveillance court",
    "reason:",
    "sensitive information",
    "tracking no",
    "washington, d.c",
}
LEADING_FURNITURE_HINTS = {
    "clerk",
    "court",
    "declassify",
    "derived",
    "docket number",
    "filed",
    "foreign intelligence",
    "reason",
    "secret",
    "tracking",
    "unclassified",
    "united states",
    "washington",
}
REFUSAL_PATTERNS = [
    re.compile(r"^\s*i(?:'| a)?m\s+sorry\b", re.IGNORECASE),
    re.compile(r"\bI can(?:not|'t)\s+(?:assist|help|process)\b", re.IGNORECASE),
    re.compile(r"\bI'm unable to process documents in that way\b", re.IGNORECASE),
    re.compile(r"\bI can help summarize\b", re.IGNORECASE),
]
PAGE_TAG_LINE_RE = re.compile(r"(?im)^\s*\[PAGE\s+\d+\]\s*$")
MONTH_MAP = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}
PAIR_METADATA_COLUMNS = [
    "collection",
    "record_id",
    "date_display",
    "source_pdf_url",
    "date_compact",
    "cia_document_id",
    "document_url",
    "meta_path",
    "title",
    "year",
    "paired_flag",
]


@dataclass
class PairVersionMeta:
    pair_label: str
    docket: str
    title: str
    version: str
    source: str
    pdf_pages: str
    document_url: str
    landing_page_url: str
    extracted_url: str
    notes: str


@dataclass
class FISCPair:
    pair_number: int
    pair_key: str
    redacted_pdf: Path
    unredacted_pdf: Path
    redacted_meta: PairVersionMeta | None
    unredacted_meta: PairVersionMeta | None


@dataclass
class ExtractionResult:
    pages: list[PageBlock]
    raw_text: str
    method: str
    page_count: int
    total_tokens: int
    substantive_pages: int
    low_content_pages: int
    usable: bool
    reason: str


def _progress(items: Iterable[Any], *, total: int | None, desc: str) -> Iterable[Any]:
    if tqdm is None:
        return items
    return tqdm(items, total=total, desc=desc, dynamic_ncols=True, leave=False, mininterval=0.2)


def _normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", str(text))
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _normalize_pair_label(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return f"{value.month}-{value.day}"
    text = str(value).strip()
    if not text:
        return None
    match = re.fullmatch(r"(\d+)\s*-\s*([12])", text)
    if match:
        return f"{int(match.group(1))}-{match.group(2)}"
    return None


def load_fisc_metadata(xlsx_path: Path) -> dict[str, PairVersionMeta]:
    if load_workbook is None:
        return {}
    if not xlsx_path.exists():
        return {}
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "FISC Pairs" not in wb.sheetnames:
        return {}
    ws = wb["FISC Pairs"]
    metadata: dict[str, PairVersionMeta] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row:
            continue
        pair_label = _normalize_pair_label(row[0])
        if not pair_label:
            continue
        width = max(10, len(row))
        padded = list(row) + [None] * max(0, width - len(row))
        metadata[f"Pair_{pair_label}"] = PairVersionMeta(
            pair_label=f"Pair_{pair_label}",
            docket=str(padded[1] or "").strip(),
            title=str(padded[2] or "").strip(),
            version=str(padded[3] or "").strip(),
            source=str(padded[4] or "").strip(),
            pdf_pages=str(padded[5] or "").strip(),
            document_url=str(padded[6] or "").strip(),
            landing_page_url=str(padded[7] or "").strip(),
            extracted_url=str(padded[8] or "").strip(),
            notes=str(padded[9] or "").strip(),
        )
    return metadata


def discover_fisc_pairs(pdf_root: Path, metadata_by_label: dict[str, PairVersionMeta]) -> list[FISCPair]:
    grouped: dict[int, dict[str, Path]] = {}
    for pdf in sorted(pdf_root.glob("Pair_*-*.pdf")):
        match = PAIR_FILENAME_RE.match(pdf.name)
        if not match:
            continue
        pair_number = int(match.group(1))
        version = match.group(2)
        grouped.setdefault(pair_number, {})[version] = pdf

    pairs: list[FISCPair] = []
    for pair_number in sorted(grouped):
        versions = grouped[pair_number]
        if "1" not in versions or "2" not in versions:
            continue
        pair_key = f"Pair_{pair_number}"
        pairs.append(
            FISCPair(
                pair_number=pair_number,
                pair_key=pair_key,
                redacted_pdf=versions["1"],
                unredacted_pdf=versions["2"],
                redacted_meta=metadata_by_label.get(f"{pair_key}-1"),
                unredacted_meta=metadata_by_label.get(f"{pair_key}-2"),
            )
        )
    return pairs


def pdf_page_count(pdf_path: Path) -> int:
    result = subprocess.run(
        ["pdfinfo", str(pdf_path)],
        capture_output=True,
        text=True,
        check=True,
    )
    match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, flags=re.MULTILINE)
    if not match:
        raise RuntimeError(f"Could not determine page count for {pdf_path}")
    return int(match.group(1))


def render_pages_to_document_text(pages: list[PageBlock]) -> str:
    width = max(2, len(str(max((page.source_page_no for page in pages), default=1))))
    lines: list[str] = []
    for page in pages:
        lines.append(f"[PAGE {page.source_page_no:0{width}d}]")
        if page.text:
            lines.append(page.text.rstrip())
    return "\n".join(lines).rstrip() + "\n"


def line_is_probable_boilerplate(line: str, repeated_lines: set[str]) -> bool:
    text = " ".join(line.split()).strip()
    if not text:
        return True
    lower = text.lower()
    if any(pattern.match(text) for pattern in FURNITURE_PATTERNS):
        return True
    if CLASSIFICATION_LINE_RE.match(text):
        return True
    if HEADER_NOISE_PREFIX_RE.match(text):
        return True
    if text in repeated_lines and any(keyword in lower for keyword in BOILERPLATE_KEYWORDS):
        return True
    if "clerk of court" in lower or "foreign intelligence surveillance court" in lower:
        return True
    if "sensitive information" in lower and len(text) <= 80:
        return True
    if "tracking no" in lower:
        return True
    if "tracking n o" in lower:
        return True
    if re.fullmatch(r"(?:[-–—_=~]+\s*){3,}", text):
        return True
    if re.fullmatch(r"(?:[A-Za-z]\d+[A-Za-z-]*[\s,]*){2,}", text):
        return True
    if "washington" in lower and "d.c" in lower and len(text) <= 40:
        return True
    return False


def line_is_probable_leading_furniture(line: str, repeated_lines: set[str]) -> bool:
    text = " ".join(line.split()).strip()
    if not text:
        return True
    if line_is_probable_boilerplate(text, repeated_lines):
        return True
    lower = text.lower()
    token_count = len(TOKEN_RE.findall(text))
    if token_count <= 8 and any(hint in lower for hint in LEADING_FURNITURE_HINTS):
        return True
    if "u.s. foreign" in lower or "surveillance court" in lower:
        return True
    return False


def clean_local_page_texts(raw_pages: list[str]) -> list[str]:
    normalized_pages = [_normalize_text(page) for page in raw_pages]
    line_counts: dict[str, int] = {}
    for page in normalized_pages:
        seen: set[str] = set()
        for raw_line in page.splitlines():
            text = " ".join(raw_line.split()).strip()
            if not text or text in seen:
                continue
            line_counts[text] = line_counts.get(text, 0) + 1
            seen.add(text)
    repeated_lines = {line for line, count in line_counts.items() if count >= 3}

    cleaned_pages: list[str] = []
    for page in normalized_pages:
        raw_lines = page.splitlines()
        start = 0
        while start < len(raw_lines):
            text = " ".join(raw_lines[start].split()).strip()
            if not text:
                start += 1
                continue
            if line_is_probable_leading_furniture(text, repeated_lines):
                start += 1
                continue
            break

        end = len(raw_lines)
        while end > start:
            text = " ".join(raw_lines[end - 1].split()).strip()
            if not text:
                end -= 1
                continue
            if line_is_probable_boilerplate(text, repeated_lines):
                end -= 1
                continue
            break

        kept_lines: list[str] = []
        for raw_line in raw_lines[start:end]:
            stripped = raw_line.strip()
            text = " ".join(stripped.split()).strip()
            if not text:
                kept_lines.append("")
                continue
            if line_is_probable_boilerplate(text, repeated_lines):
                continue
            kept_lines.append(stripped.rstrip())
        cleaned_text = "\n".join(kept_lines)
        cleaned_text = re.sub(r"\n{3,}", "\n\n", cleaned_text).strip("\n")
        cleaned_pages.append(cleaned_text)
    return cleaned_pages


def extract_pages_with_pdftotext(pdf_path: Path) -> ExtractionResult:
    page_count = pdf_page_count(pdf_path)
    result = subprocess.run(
        ["pdftotext", "-layout", str(pdf_path), "-"],
        capture_output=True,
        text=True,
        check=True,
    )
    raw_text = _normalize_text(result.stdout)
    raw_pages = raw_text.split("\f")
    if raw_pages and not raw_pages[-1].strip():
        raw_pages = raw_pages[:-1]
    if len(raw_pages) < page_count:
        raw_pages.extend([""] * (page_count - len(raw_pages)))
    elif len(raw_pages) > page_count:
        raw_pages = raw_pages[:page_count]

    cleaned_pages = clean_local_page_texts(raw_pages)
    pages = [PageBlock(source_page_no=i, source_label=f"{i:02d}", text=text) for i, text in enumerate(cleaned_pages, start=1)]
    usable, reason, total_tokens, substantive_pages, low_content_pages = assess_local_extraction(pages)
    return ExtractionResult(
        pages=pages,
        raw_text=render_pages_to_document_text(pages),
        method="pdftotext",
        page_count=page_count,
        total_tokens=total_tokens,
        substantive_pages=substantive_pages,
        low_content_pages=low_content_pages,
        usable=usable,
        reason=reason,
    )


def sanitize_openai_transcription_text(text: str) -> str:
    src = _normalize_text(text).strip()
    if src.startswith("```"):
        src = re.sub(r"^```[A-Za-z0-9_-]*\s*\n?", "", src)
        src = re.sub(r"\n?```$", "", src.strip())
    return src.strip() + "\n"


def strip_embedded_page_tags(text: str) -> str:
    src = _normalize_text(text)
    src = PAGE_TAG_LINE_RE.sub("", src)
    src = re.sub(r"\n{3,}", "\n\n", src)
    return src.strip()


def looks_like_transcription_refusal(text: str) -> bool:
    sample = _normalize_text(text).strip()
    if not sample:
        return False
    first = "\n".join(sample.splitlines()[:8]).strip()
    return any(pattern.search(first) for pattern in REFUSAL_PATTERNS)


def extract_openai_response_text(response: Any) -> str:
    text = getattr(response, "output_text", None)
    if text:
        return str(text)
    parts: list[str] = []
    for item in getattr(response, "output", []) or []:
        if getattr(item, "type", None) != "message":
            continue
        for content in getattr(item, "content", []) or []:
            if getattr(content, "type", None) == "output_text":
                parts.append(str(getattr(content, "text", "")))
    return "\n".join(part for part in parts if part).strip()


def render_pdf_pages_to_png(pdf_path: Path, out_dir: Path, *, dpi: int = 200) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    prefix = out_dir / "page"
    cmd = [
        "pdftoppm",
        "-r",
        str(dpi),
        "-png",
        str(pdf_path),
        str(prefix),
    ]
    subprocess.run(cmd, check=True, capture_output=True, text=True)
    return sorted(out_dir.glob("page-*.png"))


def transcribe_image_page_with_openai(
    client: Any,
    *,
    image_path: Path,
    model: str,
    prompt: str,
) -> str:
    payload = base64.b64encode(image_path.read_bytes()).decode("ascii")
    response = client.responses.create(
        model=model,
        input=[
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": prompt},
                    {
                        "type": "input_image",
                        "image_url": f"data:image/png;base64,{payload}",
                    },
                ],
            }
        ],
    )
    return extract_openai_response_text(response)


def extract_pages_with_openai_images(
    client: Any,
    *,
    pdf_path: Path,
    model: str,
    prompt: str,
    render_root: Path,
) -> ExtractionResult:
    image_dir = render_root / pdf_path.stem
    if image_dir.exists():
        shutil.rmtree(image_dir)
    png_paths = render_pdf_pages_to_png(pdf_path, image_dir)
    pages: list[PageBlock] = []
    raw_lines: list[str] = []
    for page_index, image_path in enumerate(png_paths, start=1):
        page_prompt = (
            prompt
            + "\n\nThis is a single page image. Return only the plain text of this page. "
            + "Do not add explanations, code fences, or any page label."
        )
        text = sanitize_openai_transcription_text(
            transcribe_image_page_with_openai(
                client,
                image_path=image_path,
                model=model,
                prompt=page_prompt,
            )
        ).strip()
        if looks_like_transcription_refusal(text):
            return ExtractionResult(
                pages=[],
                raw_text=text.rstrip() + "\n" if text else "",
                method=f"openai_image:{model}",
                page_count=0,
                total_tokens=0,
                substantive_pages=0,
                low_content_pages=0,
                usable=False,
                reason="openai_refusal",
            )
        text = strip_embedded_page_tags(text)
        pages.append(PageBlock(source_page_no=page_index, source_label=f"{page_index:02d}", text=text))
        raw_lines.append(f"[PAGE {page_index:02d}]")
        if text:
            raw_lines.append(text)
    raw_text = "\n".join(raw_lines).rstrip() + "\n"
    usable, reason, total_tokens, substantive_pages, low_content_pages = assess_local_extraction(pages)
    return ExtractionResult(
        pages=pages,
        raw_text=raw_text,
        method=f"openai_image:{model}",
        page_count=len(pages),
        total_tokens=total_tokens,
        substantive_pages=substantive_pages,
        low_content_pages=low_content_pages,
        usable=usable,
        reason=reason,
    )


def assess_local_extraction(pages: list[PageBlock]) -> tuple[bool, str, int, int, int]:
    nonempty_pages = 0
    substantive_pages = 0
    low_content_pages = 0
    total_tokens = 0
    all_furniture = True

    for page in pages:
        page_text = page.text.strip()
        if not page_text:
            continue
        nonempty_pages += 1
        lines = [line.strip() for line in page_text.splitlines() if line.strip()]
        content_lines = [line for line in lines if not any(pattern.match(line) for pattern in FURNITURE_PATTERNS)]
        content_text = " ".join(content_lines).strip()
        token_count = len(TOKEN_RE.findall(content_text))
        alpha_count = sum(ch.isalpha() for ch in content_text)
        total_tokens += token_count
        if content_lines:
            all_furniture = False
        if token_count <= 3 or alpha_count < 20:
            low_content_pages += 1
        if token_count >= 25 or alpha_count >= 120:
            substantive_pages += 1

    if nonempty_pages == 0:
        return False, "empty_extraction", total_tokens, substantive_pages, low_content_pages
    if all_furniture:
        return False, "page_furniture_only", total_tokens, substantive_pages, low_content_pages
    if substantive_pages == 0:
        return False, "no_substantive_pages", total_tokens, substantive_pages, low_content_pages
    if low_content_pages == nonempty_pages and substantive_pages <= 1:
        return False, "all_pages_low_content", total_tokens, substantive_pages, low_content_pages
    return True, "usable_local_text", total_tokens, substantive_pages, low_content_pages


def maybe_extract_pdf(
    *,
    pdf_path: Path,
    openai_model: str,
    prompt: str,
    cached_client: dict[str, Any],
    image_render_root: Path,
) -> ExtractionResult:
    if "client" not in cached_client:
        cached_client["client"] = ensure_openai_client()
    return extract_pages_with_openai_images(
        cached_client["client"],
        pdf_path=pdf_path,
        model=openai_model,
        prompt=prompt,
        render_root=image_render_root,
    )


def infer_pair_date(text: str, title: str) -> tuple[str, str]:
    src = _normalize_text(text[:12000] + "\n" + title)

    match = FULL_DATE_RE.search(src)
    if match:
        month = MONTH_MAP[match.group(1).lower()]
        day = int(match.group(2))
        year = int(match.group(3))
        dt = datetime(year, month, day)
        return dt.strftime("%Y-%m-%d"), dt.strftime("%Y%m%d")

    match = DAY_MONTH_YEAR_RE.search(src)
    if match:
        day = int(match.group(1))
        month = MONTH_MAP[match.group(2).lower()]
        year = int(match.group(3))
        dt = datetime(year, month, day)
        return dt.strftime("%Y-%m-%d"), dt.strftime("%Y%m%d")

    match = YEAR_MONTH_RE.search(src)
    if match:
        year = int(match.group(1))
        month = MONTH_MAP[match.group(2).lower()]
        dt = datetime(year, month, 1)
        return dt.strftime("%Y-%m-%d"), dt.strftime("%Y%m%d")

    match = MONTH_YEAR_RE.search(src)
    if match:
        month = MONTH_MAP[match.group(1).lower()]
        year = int(match.group(2))
        dt = datetime(year, month, 1)
        return dt.strftime("%Y-%m-%d"), dt.strftime("%Y%m%d")

    return "", ""


def build_pair_metadata(pair: FISCPair, raw_unredacted_text: str, workbook_path: Path) -> dict[str, Any]:
    chosen = pair.unredacted_meta or pair.redacted_meta
    title = chosen.title if chosen and chosen.title else pair.pair_key
    collection = "FISC Carter Page"
    docket = chosen.docket if chosen and chosen.docket else pair.pair_key
    document_url = chosen.document_url if chosen and chosen.document_url else ""
    source_pdf_url = chosen.extracted_url if chosen and chosen.extracted_url else document_url
    date_display, date_compact = infer_pair_date(raw_unredacted_text, title)
    year = date_display[:4] if date_display else ""

    csv_row = [
        collection,
        pair.pair_key,
        date_display,
        source_pdf_url,
        date_compact,
        docket,
        document_url,
        str(workbook_path),
        title,
        year,
        "1",
    ]
    while len(csv_row) < len(PAIR_METADATA_COLUMNS):
        csv_row.append("")

    return {
        "pair_key": pair.pair_key,
        "model": "openai_pdf_pair_pipeline",
        "csv_row": csv_row,
        "redacted_pdf": str(pair.redacted_pdf),
        "unredacted_pdf": str(pair.unredacted_pdf),
        "redacted_metadata": _json_safe(pair.redacted_meta.__dict__) if pair.redacted_meta else {},
        "unredacted_metadata": _json_safe(pair.unredacted_meta.__dict__) if pair.unredacted_meta else {},
    }


def write_pair_outputs(
    *,
    pair: FISCPair,
    out_root: Path,
    redacted_result: ExtractionResult,
    unredacted_result: ExtractionResult,
    workbook_path: Path,
) -> dict[str, Any]:
    pair_dir = out_root / pair.pair_key
    document_dir = pair_dir / "document"
    pages_dir = pair_dir / "pages"
    redacted_pages_dir = pages_dir / "redacted"
    unredacted_pages_dir = pages_dir / "unredacted"
    source_pdf_dir = pair_dir / "source_pdfs"

    document_dir.mkdir(parents=True, exist_ok=True)
    redacted_pages_dir.mkdir(parents=True, exist_ok=True)
    unredacted_pages_dir.mkdir(parents=True, exist_ok=True)
    source_pdf_dir.mkdir(parents=True, exist_ok=True)

    shutil.copy2(pair.redacted_pdf, source_pdf_dir / pair.redacted_pdf.name)
    shutil.copy2(pair.unredacted_pdf, source_pdf_dir / pair.unredacted_pdf.name)

    (document_dir / "redacted.transcription.txt").write_text(redacted_result.raw_text, encoding="utf-8")
    (document_dir / "unredacted.transcription.txt").write_text(unredacted_result.raw_text, encoding="utf-8")

    red_rows = save_split_pages(redacted_pages_dir, redacted_result.pages)
    unred_rows = save_split_pages(unredacted_pages_dir, unredacted_result.pages)
    (pages_dir / "redacted.pages.json").write_text(json.dumps(red_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    (pages_dir / "unredacted.pages.json").write_text(json.dumps(unred_rows, ensure_ascii=False, indent=2), encoding="utf-8")

    metadata = build_pair_metadata(pair, unredacted_result.raw_text, workbook_path)
    metadata["redacted_extraction"] = {
        "method": redacted_result.method,
        "usable": redacted_result.usable,
        "reason": redacted_result.reason,
        "page_count": redacted_result.page_count,
        "total_tokens": redacted_result.total_tokens,
        "substantive_pages": redacted_result.substantive_pages,
        "low_content_pages": redacted_result.low_content_pages,
    }
    metadata["unredacted_extraction"] = {
        "method": unredacted_result.method,
        "usable": unredacted_result.usable,
        "reason": unredacted_result.reason,
        "page_count": unredacted_result.page_count,
        "total_tokens": unredacted_result.total_tokens,
        "substantive_pages": unredacted_result.substantive_pages,
        "low_content_pages": unredacted_result.low_content_pages,
    }
    (document_dir / "pair_metadata.json").write_text(json.dumps(_json_safe(metadata), ensure_ascii=False, indent=2), encoding="utf-8")

    alignment_rows = global_align_pages(redacted_result.pages, unredacted_result.pages)
    pair_ref = PDFPair(
        pair_key=pair.pair_key,
        unredacted_numeric_id=str(pair.pair_number),
        redacted_doc_id=str(pair.pair_number),
        row=[""] * 11,
        row_index_1based=pair.pair_number,
        redacted_pdf=pair.redacted_pdf,
        unredacted_pdf=pair.unredacted_pdf,
    )
    build_difference_artifacts(
        pair=pair_ref,
        red_pages=redacted_result.pages,
        unred_pages=unredacted_result.pages,
        alignment_rows=alignment_rows,
        out_root=out_root,
    )
    return {
        "pair_key": pair.pair_key,
        "redacted_pdf": pair.redacted_pdf.name,
        "unredacted_pdf": pair.unredacted_pdf.name,
        "redacted_method": redacted_result.method,
        "unredacted_method": unredacted_result.method,
        "redacted_reason": redacted_result.reason,
        "unredacted_reason": unredacted_result.reason,
        "redacted_tokens": redacted_result.total_tokens,
        "unredacted_tokens": unredacted_result.total_tokens,
        "aligned_pages": len(alignment_rows),
    }


def select_pairs(
    pairs: list[FISCPair],
    *,
    pair_key: str | None,
    max_pairs: int | None,
) -> list[FISCPair]:
    selected = pairs
    if pair_key:
        selected = [pair for pair in selected if pair.pair_key == pair_key]
    if max_pairs is not None:
        selected = selected[: max(0, int(max_pairs))]
    return selected


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare FISC PDF pairs into a packaged prepared corpus root.")
    parser.add_argument(
        "--pdf-root",
        default=str(PACKAGE_ROOT / "data" / "raw" / "fisc"),
        help="Directory containing Pair_*-*.pdf files.",
    )
    parser.add_argument(
        "--xlsx-path",
        default=str(PACKAGE_ROOT / "data" / "raw" / "fisc" / "FISC Carter Page redacted pairs.xlsx"),
        help="Optional workbook with pair metadata.",
    )
    parser.add_argument(
        "--transcription-root",
        default=str(PACKAGE_ROOT / "data" / "intermediate" / "fisc" / "transcription_staging"),
        help="Staging root for extracted pair transcriptions.",
    )
    parser.add_argument(
        "--prepared-root",
        default=str(PACKAGE_ROOT / "data" / "prepared" / "fisc"),
        help="Final prepared root for the packaged KG pipeline.",
    )
    parser.add_argument("--pair-key", default=None, help="Run exactly one pair key such as Pair_1")
    parser.add_argument("--max-pairs", type=int, default=None, help="Optional cap for smoke tests")
    parser.add_argument("--openai-model", default="gpt-4o-mini", help="OpenAI model for page transcription")
    parser.add_argument("--skip-postprocess", action="store_true", help="Only build the staging transcription root")
    parser.add_argument(
        "--cleanup-staging",
        action="store_true",
        help="Delete the transcription staging root after successful postprocess.",
    )
    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing pair directories")
    parser.add_argument("--stop-on-error", action="store_true", help="Abort immediately if one pair fails")
    parser.add_argument("--image-render-root", default="", help="Optional temp directory for rendered PNG pages used in OpenAI image OCR")
    args = parser.parse_args()

    pdf_root = Path(args.pdf_root).expanduser().resolve()
    xlsx_path = Path(args.xlsx_path).expanduser().resolve()
    transcription_root = Path(args.transcription_root).expanduser().resolve()
    prepared_root = Path(args.prepared_root).expanduser().resolve()
    transcription_root.mkdir(parents=True, exist_ok=True)

    metadata_by_label = load_fisc_metadata(xlsx_path)
    pairs = discover_fisc_pairs(pdf_root, metadata_by_label)
    selected_pairs = select_pairs(pairs, pair_key=args.pair_key, max_pairs=args.max_pairs)
    if not selected_pairs:
        raise SystemExit("No FISC PDF pairs selected.")

    cached_client: dict[str, Any] = {}
    manifest_rows: list[dict[str, Any]] = []
    manifest_path = transcription_root / "prepare_fisc_manifest.json"
    image_render_root = Path(args.image_render_root) if args.image_render_root else (transcription_root / "_rendered_page_images")

    def _write_manifest() -> None:
        manifest_path.write_text(
            json.dumps(_json_safe(manifest_rows), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    for pair in _progress(selected_pairs, total=len(selected_pairs), desc="Prepare FISC pairs"):
        pair_dir = transcription_root / pair.pair_key
        try:
            if pair_dir.exists() and args.overwrite:
                shutil.rmtree(pair_dir)
            elif pair_dir.exists() and not args.overwrite:
                manifest_rows.append({"pair_key": pair.pair_key, "status": "skipped_existing"})
                _write_manifest()
                continue

            redacted_result = maybe_extract_pdf(
                pdf_path=pair.redacted_pdf,
                openai_model=args.openai_model,
                prompt=FISC_TRANSCRIPTION_PROMPT,
                cached_client=cached_client,
                image_render_root=image_render_root,
            )
            unredacted_result = maybe_extract_pdf(
                pdf_path=pair.unredacted_pdf,
                openai_model=args.openai_model,
                prompt=FISC_TRANSCRIPTION_PROMPT,
                cached_client=cached_client,
                image_render_root=image_render_root,
            )
            row = write_pair_outputs(
                pair=pair,
                out_root=transcription_root,
                redacted_result=redacted_result,
                unredacted_result=unredacted_result,
                workbook_path=xlsx_path,
            )
            row["status"] = "prepared"
            manifest_rows.append(row)
            _write_manifest()
        except Exception as exc:
            manifest_rows.append(
                {
                    "pair_key": pair.pair_key,
                    "status": "error",
                    "error": str(exc),
                    "redacted_pdf": str(pair.redacted_pdf),
                    "unredacted_pdf": str(pair.unredacted_pdf),
                }
            )
            _write_manifest()
            print(f"[WARN] failed on {pair.pair_key}: {exc}", flush=True)
            if args.stop_on_error:
                raise

    if args.skip_postprocess:
        return

    run_postprocess(
        in_root=transcription_root,
        out_root=prepared_root,
        pair_key=args.pair_key,
        max_pairs=args.max_pairs,
        copy_source_pdfs=True,
        min_tokens=1,
        full_page_coverage=0.85,
        low_alignment_score=0.2,
        min_full_page_tokens=30,
        ocr_similarity_threshold=0.88,
        short_ocr_similarity_threshold=0.94,
        alternate_near_full_page_coverage=0.72,
        progress_every=1,
    )

    if args.cleanup_staging and transcription_root.exists():
        shutil.rmtree(transcription_root)


if __name__ == "__main__":
    main()
